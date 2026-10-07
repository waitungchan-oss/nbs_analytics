from datetime import datetime

import pandas as pd
import pytest
from openpyxl import load_workbook


def _detail_frame(prefix, receipt_times, amounts, *, source_tag):
    rows = []
    for index, (receipt_time, amount) in enumerate(zip(receipt_times, amounts), start=1):
        rows.append(
            {
                "來源單據號": f"{prefix}{index}",
                "統一日期": "2025-01-10",
                "交易時間": "2025-01-10 09:00:00",
                "收款時間": receipt_time,
                "銷售點": "銅鑼灣分社",
                "銷售員": "測試銷售員",
                "收款原幣金額": amount,
                "收款類型": "正常收款",
                "收款方式": "現金",
                "團負責人部門": "",
                "團名稱": "測試產品",
                "來源報表標籤": source_tag,
                "行程天數": 1,
                "數量": 1,
            }
        )
    return pd.DataFrame(rows)


def test_exported_receipt_times_are_excel_datetimes_on_all_detail_sheets():
    import pipeline

    tour = _detail_frame(
        "T",
        ["2025-01-10 09:15:00", pd.Timestamp("2025-01-11 10:20:00"), ""],
        [100, 200, 0],
        source_tag="旅行團",
    )
    others = _detail_frame(
        "O",
        [pd.Timestamp("2025-01-10 11:25:00"), "2025-01-12 12:30:00", None],
        [25, 75, 0],
        source_tag="門券all",
    )

    workbook_bytes, _, _ = pipeline.build_dashboard_data(
        tour,
        others,
        {"33": "銅鑼灣分社", "225": "營銷運營中心-專職銷售組"},
        ["銅鑼灣分社"],
        [],
        [],
    )

    workbook = load_workbook(workbook_bytes, data_only=True)
    assert len(workbook.sheetnames) == 18

    for sheet_name in (
        "總表_多表匹配完成",
        "旅行團_匹配成功",
        "其它_未匹配_包含其它業務",
    ):
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        receipt_column = headers.index("收款時間") + 1
        receipt_cells = [sheet.cell(row=row, column=receipt_column) for row in range(2, sheet.max_row + 1)]
        populated = [cell for cell in receipt_cells if cell.value is not None]

        assert populated
        assert all(isinstance(cell.value, datetime) for cell in populated)
        assert all(cell.number_format.lower() == "yyyy-mm-dd hh:mm:ss" for cell in populated)
        assert any(cell.value is None for cell in receipt_cells)

    def amount_total(sheet_name):
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        amount_column = headers.index("收款原幣金額") + 1
        return sum(float(sheet.cell(row=row, column=amount_column).value or 0) for row in range(2, sheet.max_row + 1))

    assert amount_total("旅行團_匹配成功") == 300
    assert amount_total("其它_未匹配_包含其它業務") == 100
    assert amount_total("總表_多表匹配完成") == 400


def test_export_rejects_unparseable_nonblank_receipt_time():
    import pipeline

    tour = _detail_frame("T", ["not-a-date"], [100], source_tag="旅行團")
    others = _detail_frame("O", ["2025-01-10 11:25:00"], [25], source_tag="門券all")

    with pytest.raises(ValueError, match="總表_多表匹配完成.*收款時間.*1 筆無法解析"):
        pipeline.build_dashboard_data(
            tour,
            others,
            {"33": "銅鑼灣分社", "225": "營銷運營中心-專職銷售組"},
            ["銅鑼灣分社"],
            [],
            [],
        )
