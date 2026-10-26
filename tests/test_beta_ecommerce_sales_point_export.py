import pandas as pd
import pytest
from pathlib import Path


E_COMMERCE = "市場及電商部-電子商務組"
BETA_SALES_POINTS = (
    E_COMMERCE,
    "營銷運營中心-同業銷售部",
    "市場及商務部-電子商務組",
    "網銷組 i6",
)
BETA_LABEL = "市場電商及同業"


def _row(order_id, sales_point, salesperson, department=""):
    return {
        "來源單據號": order_id,
        "統一日期": "2026-05-01",
        "銷售點": sales_point,
        "銷售員": salesperson,
        "收款原幣金額": 100,
        "收款類型": "正常收款",
        "收款方式": "現金",
        "團負責人部門": department,
        "交易時間": "2026-05-01",
        "行程天數": 3,
        "數量": 2,
        "團名稱": "測試行程",
        "來源報表標籤": "旅行團",
    }


def tour_frame():
    return pd.DataFrame(
        [
            _row("T001", E_COMMERCE, "Alice"),
            _row("T002", E_COMMERCE, ""),
            _row("T003", "銅鑼灣分社", "Wrong", E_COMMERCE),
        ]
    )


def others_frame():
    return pd.DataFrame([_row("O001", E_COMMERCE, "Bob"), _row("O002", "銅鑼灣分社", "Other")])


def empty_frame():
    return tour_frame().iloc[0:0].copy()


def branch_mapping():
    return {"I6": E_COMMERCE, "225": "營銷運營中心-專職銷售組"}


def beta_category_frames():
    salespeople = [
        "011185 洪淑芬",
        "CTSIT 杨慧如",
        "debby 卢淑贞",
        "004014 杨慧如",
        "New Ecommerce Rep",
        " NaN ",
    ]
    tours = []
    tickets = []
    for index, salesperson in enumerate(salespeople, start=1):
        tours.append(_row(f"BT{index:03}", E_COMMERCE, salesperson))
        ticket = _row(f"BO{index:03}", E_COMMERCE, salesperson)
        ticket["來源報表標籤"] = "門券"
        tickets.append(ticket)
    return pd.DataFrame(tours), pd.DataFrame(tickets)


def test_beta_selects_sales_point_and_keeps_all_salespeople():
    import pipeline

    beta_tour, beta_others = pipeline._select_sales_point_frames(
        tour_frame(), others_frame(), sales_point=E_COMMERCE
    )

    assert set(beta_tour["銷售員"]) == {"Alice", "未指定"}
    assert set(beta_others["銷售員"]) == {"Bob"}
    assert set(beta_tour["銷售點"]) == {E_COMMERCE}


def test_beta_does_not_use_department_as_fallback():
    import pipeline

    beta_tour, _ = pipeline._select_sales_point_frames(
        tour_frame(), others_frame(), sales_point=E_COMMERCE
    )

    assert "Wrong" not in set(beta_tour["銷售員"])
    assert len(beta_tour) == 2


def test_beta_missing_salesperson_column_becomes_unspecified():
    import pipeline

    tours = tour_frame().drop(columns=["銷售員"])
    beta_tour, _ = pipeline._select_sales_point_frames(
        tours, empty_frame(), sales_point=E_COMMERCE
    )

    assert set(beta_tour["銷售員"]) == {"未指定"}


def test_beta_does_not_filter_by_legacy_sales_rep_list():
    import pipeline

    frames = tour_frame().assign(銷售員="New Ecommerce Rep")
    beta_tour, _ = pipeline._select_sales_point_frames(frames, others_frame(), sales_point=E_COMMERCE)

    assert "New Ecommerce Rep" in set(beta_tour["銷售員"])


def test_beta_selection_does_not_mutate_sources():
    import pipeline

    tour = tour_frame()
    others = others_frame()
    before_tour = tour.copy(deep=True)
    before_others = others.copy(deep=True)

    pipeline._select_sales_point_frames(tour, others, sales_point=E_COMMERCE)

    pd.testing.assert_frame_equal(tour, before_tour)
    pd.testing.assert_frame_equal(others, before_others)


def test_beta_workbook_replaces_specialist_sheets_with_all_ecommerce_salespeople():
    import pipeline

    _, _, facts = pipeline.build_dashboard_data(
        tour_frame(),
        others_frame(),
        branch_mapping(),
        [],
        [],
        ["Legacy Rep"],
        make_workbook=False,
        return_facts=True,
        beta_sales_point=BETA_SALES_POINTS,
    )

    assert f"{BETA_LABEL}_經營統計" in facts
    assert set(facts[f"{BETA_LABEL}_經營統計"]["種類"]) == {"客戶服務部"}
    assert set(facts[f"{BETA_LABEL}_經營統計"]["文本"]) >= {"Alice", "未指定"}
    assert set(facts[f"{BETA_LABEL}_每天旅行團交易人數"]["銷售員"]) >= {"Alice", "未指定"}
    assert "Legacy Rep" not in set(facts[f"{BETA_LABEL}_每天旅行團交易人數"]["銷售員"])


def test_beta_combines_all_four_sales_points_into_six_named_sheets():
    import pipeline

    combined_tour = pd.concat(
        [
            tour_frame(),
            pd.DataFrame([_row("T004", "市場及商務部-電子商務組", "Business Rep")]),
        ],
        ignore_index=True,
    )
    combined_others = pd.concat(
        [
            others_frame(),
            pd.DataFrame(
                [
                    _row("O003", "營銷運營中心-同業銷售部", "Industry Rep"),
                    _row("O004", "網銷組 i6", "Network Rep"),
                ]
            ),
        ],
        ignore_index=True,
    )

    _, _, facts = pipeline.build_dashboard_data(
        combined_tour,
        combined_others,
        branch_mapping(),
        [],
        [],
        ["Legacy Rep"],
        make_workbook=False,
        return_facts=True,
        beta_sales_point=BETA_SALES_POINTS,
    )

    expected_sheets = {
        f"{BETA_LABEL}_經營統計",
        f"{BETA_LABEL}_旅行團統計",
        f"{BETA_LABEL}_票務總計",
        f"{BETA_LABEL}_每天旅行團交易人數",
        f"{BETA_LABEL}_每天票務交易數量",
        f"{BETA_LABEL}_線路種類每天統計",
    }
    assert expected_sheets <= set(facts)
    assert set(facts[f"{BETA_LABEL}_經營統計"]["文本"]) >= {
        "Alice",
        "Bob",
        "Business Rep",
        "Industry Rep",
        "Network Rep",
        "未指定",
    }


def test_beta_export_normalizes_mixed_transaction_time_types_before_sorting():
    import pipeline

    tour = tour_frame()
    tour.loc[0, "交易時間"] = pd.Timestamp("2026-05-01 09:00:00")
    others = others_frame()
    others.loc[0, "交易時間"] = pd.Timestamp("2026-05-01 10:00:00")

    _, _, facts = pipeline.build_dashboard_data(
        tour,
        others,
        branch_mapping(),
        [],
        [],
        ["Legacy Rep"],
        make_workbook=False,
        return_facts=True,
        beta_sales_point=BETA_SALES_POINTS,
    )

    assert not facts[f"{BETA_LABEL}_旅行團統計"].empty


def test_beta_empty_sales_point_keeps_schema_without_legacy_fallback():
    import pipeline

    _, _, facts = pipeline.build_dashboard_data(
        empty_frame(),
        empty_frame(),
        branch_mapping(),
        [],
        [],
        ["Legacy Rep"],
        make_workbook=False,
        return_facts=True,
        beta_sales_point=BETA_SALES_POINTS,
    )

    expected_schemas = {
        f"{BETA_LABEL}_旅行團統計": [
            "銷售員", "所屬種類", "文本", "天數", "日期", "月份", "交易人數"
        ],
        f"{BETA_LABEL}_票務總計": [
            "銷售員", "所屬種類", "文本", "日期", "月份", "交易數量"
        ],
        f"{BETA_LABEL}_每天旅行團交易人數": [
            "銷售員", "所屬種類", "日期", "月份", "旅行團交易人數", "郵輪交易人數"
        ],
        f"{BETA_LABEL}_每天票務交易數量": [
            "銷售員", "所屬種類", "種類", "日期", "月份", "交易數量"
        ],
        f"{BETA_LABEL}_線路種類每天統計": [
            "銷售員", "所屬種類", "線路種類", "日子", "月份", "交易人數", "交易金額"
        ],
    }
    for sheet_name, columns in expected_schemas.items():
        assert facts[sheet_name].empty
        assert list(facts[sheet_name].columns) == columns


def test_beta_market_summary_classifies_named_salespeople_and_defaults_others():
    import pipeline

    tours, tickets = beta_category_frames()
    _, _, facts = pipeline.build_dashboard_data(
        tours,
        tickets,
        branch_mapping(),
        [],
        [],
        ["Legacy Rep"],
        make_workbook=False,
        return_facts=True,
        beta_sales_point=BETA_SALES_POINTS,
    )

    summary = facts[f"{BETA_LABEL}_經營統計"]
    categories = dict(zip(summary["文本"], summary["種類"]))
    assert categories == {
        "011185 洪淑芬": "商務同業組",
        "CTSIT 杨慧如": "商務同業組",
        "debby 卢淑贞": "商務同業組",
        "004014 杨慧如": "商務同業組",
        "New Ecommerce Rep": "客戶服務部",
        "未指定": "客戶服務部",
    }


def test_beta_seller_normalization_does_not_change_matched_total_detail():
    import pipeline

    _, _, facts = pipeline.build_dashboard_data(
        tour_frame(),
        others_frame(),
        branch_mapping(),
        [],
        [],
        ["Legacy Rep"],
        make_workbook=False,
        return_facts=True,
        beta_sales_point=BETA_SALES_POINTS,
    )

    total = facts["總表_多表匹配完成"]
    blank_seller = total.loc[total["來源單據號"].eq("T002"), "銷售員"]
    assert len(blank_seller) == 1
    assert blank_seller.iloc[0] == ""


def test_beta_detail_sheets_split_by_salesperson_and_reconcile_totals():
    import pipeline

    tours, tickets = beta_category_frames()
    _, _, facts = pipeline.build_dashboard_data(
        tours,
        tickets,
        branch_mapping(),
        [],
        [],
        ["Legacy Rep"],
        make_workbook=False,
        return_facts=True,
        beta_sales_point=BETA_SALES_POINTS,
    )

    expected_salespeople = {
        "011185 洪淑芬",
        "CTSIT 杨慧如",
        "debby 卢淑贞",
        "004014 杨慧如",
        "New Ecommerce Rep",
        "未指定",
    }
    sheet_specs = {
        f"{BETA_LABEL}_旅行團統計": ("交易人數", 12),
        f"{BETA_LABEL}_票務總計": ("交易數量", 12),
        f"{BETA_LABEL}_每天旅行團交易人數": ("旅行團交易人數", 12),
        f"{BETA_LABEL}_每天票務交易數量": ("交易數量", 12),
        f"{BETA_LABEL}_線路種類每天統計": ("交易人數", 12),
    }
    for sheet_name, (measure, expected_total) in sheet_specs.items():
        sheet = facts[sheet_name]
        assert {"銷售員", "所屬種類"} <= set(sheet.columns), sheet_name
        assert set(sheet["銷售員"]) == expected_salespeople, sheet_name
        assert int(sheet[measure].sum()) == expected_total, sheet_name
        assert set(sheet["所屬種類"]) == {"商務同業組", "客戶服務部"}, sheet_name
        assert set(
            sheet.loc[sheet["銷售員"].eq("004014 杨慧如"), "所屬種類"]
        ) == {"商務同業組"}, sheet_name

    assert "種類" in facts[f"{BETA_LABEL}_每天票務交易數量"].columns
    assert set(facts[f"{BETA_LABEL}_每天票務交易數量"]["種類"]) == {"其它門券"}
    assert int(facts[f"{BETA_LABEL}_線路種類每天統計"]["交易金額"].sum()) == 600


def test_beta_tour_detail_sheets_exclude_rows_without_a_reportable_date():
    import pipeline

    tours = pd.DataFrame([_row("BT-NO-DATE", E_COMMERCE, "Alice")])
    tickets = pd.DataFrame([_row("BO-NO-DATE", E_COMMERCE, "Alice")])
    tickets["來源報表標籤"] = "門券"
    for frame in (tours, tickets):
        frame["統一日期"] = None
        frame["交易時間"] = None
    _, _, facts = pipeline.build_dashboard_data(
        tours,
        tickets,
        branch_mapping(),
        [],
        [],
        [],
        make_workbook=False,
        return_facts=True,
        beta_sales_point=BETA_SALES_POINTS,
    )

    assert facts[f"{BETA_LABEL}_旅行團統計"].empty
    assert facts[f"{BETA_LABEL}_每天旅行團交易人數"].empty
    assert facts[f"{BETA_LABEL}_票務總計"].empty
    assert facts[f"{BETA_LABEL}_每天票務交易數量"].empty
    assert facts[f"{BETA_LABEL}_線路種類每天統計"].empty


def test_beta_workbook_writes_salesperson_and_category_headers():
    import pipeline

    tours, tickets = beta_category_frames()
    workbook_bytes, _, _ = pipeline.build_dashboard_data(
        tours,
        tickets,
        branch_mapping(),
        [],
        [],
        ["Legacy Rep"],
        beta_sales_point=BETA_SALES_POINTS,
    )

    workbook = pd.ExcelFile(workbook_bytes)
    expected_headers = {
        f"{BETA_LABEL}_經營統計": {"文本", "種類"},
        f"{BETA_LABEL}_旅行團統計": {"銷售員", "所屬種類"},
        f"{BETA_LABEL}_票務總計": {"銷售員", "所屬種類"},
        f"{BETA_LABEL}_每天旅行團交易人數": {"銷售員", "所屬種類"},
        f"{BETA_LABEL}_每天票務交易數量": {"銷售員", "所屬種類", "種類"},
        f"{BETA_LABEL}_線路種類每天統計": {"銷售員", "所屬種類"},
    }
    assert set(expected_headers) <= set(workbook.sheet_names)
    for sheet_name, columns in expected_headers.items():
        assert columns <= set(pd.read_excel(workbook, sheet_name=sheet_name, nrows=0).columns)


def test_formal_dashboard_excludes_beta_sales_point_from_branch_aggregation():
    import pipeline

    _, _, facts = pipeline.build_dashboard_data(
        tour_frame(),
        others_frame(),
        branch_mapping(),
        [],
        [],
        ["Legacy Rep"],
        make_workbook=False,
        return_facts=True,
    )

    formal_branch_rows = facts["分社經營統計"]
    assert not formal_branch_rows["文本"].astype(str).str.contains(E_COMMERCE, regex=False).any()


def test_beta_export_has_distinct_variant_identity_and_artifacts():
    import app_workflows

    payload = app_workflows._compute_beta_export_workbooks(tour_frame(), others_frame())

    assert payload["export_variant"] == "beta_market_ecommerce_peer_sales_points_v5"
    assert payload["sales_point_filter"] == list(BETA_SALES_POINTS)
    assert {"ex_beta", "ex_no_writeoff_beta", "ex_no_writeoff_refund_transfer_beta"} <= set(payload)
    assert all(payload[key] for key in ("ex_beta", "ex_no_writeoff_beta", "ex_no_writeoff_refund_transfer_beta"))


def test_beta_workbooks_add_branch_salesperson_sheet_with_scope_reconciliation():
    from io import BytesIO

    from openpyxl import load_workbook

    import app_workflows

    tours = pd.DataFrame(
        [
            _row("BT001", "銅鑼灣分社", "Alice"),
            _row("BT002", "銅鑼灣分社", "Ben", department="郵輪部"),
            _row("BT999", "銅鑼灣分社", "Alice"),
        ]
    )
    tours.loc[tours["來源單據號"].eq("BT999"), "收款類型"] = "掛賬核銷"
    others = pd.DataFrame(
        [
            _row("BO001", "銅鑼灣分社", "Alice"),
            _row("BO002", "銅鑼灣分社", "Ben"),
            _row("BO999", "銅鑼灣分社", "Ben"),
        ]
    )
    others["來源報表標籤"] = "門券all"
    others.loc[others["來源單據號"].eq("BO999"), "收款方式"] = "TT 退款轉團款"

    payload = app_workflows._compute_beta_export_workbooks(
        tours,
        others,
        rules=(
            {"1A": "銅鑼灣分社", "225": "專職銷售組"},
            ["銅鑼灣分社"],
            ["郵輪部"],
            ["Alice", "Ben"],
            [],
        ),
    )
    expected_metrics = {
        "ex_beta": ((200, 100, 300), 6, 6),
        "ex_no_writeoff_beta": ((100, 100, 300), 4, 6),
        "ex_no_writeoff_refund_transfer_beta": ((100, 100, 200), 4, 4),
    }
    expected_headers = (
        "文本",
        "單選",
        "銷售員",
        "日期",
        "月份",
        "旅行團",
        "郵輪",
        "票務",
        "旅行團交易人數",
        "票務交易數量",
    )

    for key, (
        expected_amounts,
        expected_tour_people,
        expected_ticket_quantity,
    ) in expected_metrics.items():
        workbook = load_workbook(BytesIO(payload[key]), data_only=True, read_only=True)
        assert "分社經營統計_含銷售員" in workbook.sheetnames
        assert "分社經營統計" in workbook.sheetnames
        assert f"{BETA_LABEL}_經營統計" in workbook.sheetnames

        detailed_rows = list(
            workbook["分社經營統計_含銷售員"].iter_rows(values_only=True)
        )
        branch_rows = list(workbook["分社經營統計"].iter_rows(values_only=True))
        tour_people_rows = list(
            workbook["分社每天旅行團交易人數"].iter_rows(values_only=True)
        )
        ticket_quantity_rows = list(
            workbook["分社每天票務交易數量"].iter_rows(values_only=True)
        )
        assert detailed_rows[0] == expected_headers

        def sum_columns(rows, index):
            return sum(float(row[index] or 0) for row in rows[1:])

        assert workbook.sheetnames.index("分社經營統計_含銷售員") == (
            workbook.sheetnames.index("分社經營統計") + 1
        )
        detailed_amounts = tuple(
            sum_columns(detailed_rows, index) for index in (5, 6, 7)
        )
        branch_amounts = tuple(
            sum_columns(branch_rows, index) for index in (4, 5, 6)
        )
        assert detailed_amounts == expected_amounts
        assert detailed_amounts == branch_amounts
        assert sum_columns(detailed_rows, 8) == expected_tour_people
        assert (
            sum_columns(tour_people_rows, 3) + sum_columns(tour_people_rows, 4)
            == expected_tour_people
        )
        assert sum_columns(detailed_rows, 9) == expected_ticket_quantity
        assert sum_columns(ticket_quantity_rows, 4) == expected_ticket_quantity
        assert {row[2] for row in detailed_rows[1:]} == {"Alice", "Ben"}
        workbook.close()


def test_beta_branch_salesperson_tour_people_uses_global_order_deduplication():
    import pipeline

    first = _row("DUP001", E_COMMERCE, "Alice")
    first["交易時間"] = "2026-05-01 09:00:00"
    first["數量"] = 2
    second = _row("DUP001", "銅鑼灣分社", "Ben")
    second["交易時間"] = "2026-05-01 10:00:00"
    second["數量"] = 5
    retained_branch_row = _row("DUP002", "銅鑼灣分社", "Alice")
    retained_branch_row["交易時間"] = "2026-05-01 09:00:00"
    retained_branch_row["數量"] = 2
    later_beta_row = _row("DUP002", E_COMMERCE, "Ben")
    later_beta_row["交易時間"] = "2026-05-01 10:00:00"
    later_beta_row["數量"] = 5

    _, _, facts = pipeline.build_dashboard_data_excluding_receipt_types(
        pd.DataFrame([first, second, retained_branch_row, later_beta_row]),
        empty_frame(),
        {"1A": "銅鑼灣分社", "I6": E_COMMERCE},
        ["銅鑼灣分社"],
        [],
        [],
        [],
        make_workbook=False,
        include_branch_salesperson_sheet=True,
        return_facts=True,
        beta_sales_point=BETA_SALES_POINTS,
    )

    detail = facts["分社經營統計_含銷售員"]
    official_daily = facts["分社每天旅行團交易人數"]
    assert detail["旅行團交易人數"].sum() == 2
    assert official_daily["交易人數"].sum() + official_daily["郵輪交易人數"].sum() == 2
    assert detail.loc[detail["銷售員"].eq("Ben"), "旅行團交易人數"].sum() == 0
    assert detail.loc[
        detail["銷售員"].eq("Alice") & detail["文本"].str.endswith("銅鑼灣分社"),
        "旅行團交易人數",
    ].sum() == 2


def test_beta_export_cache_requires_current_variant_and_all_three_workbooks():
    import app_workflows

    keys = ("ex_beta", "ex_no_writeoff_beta", "ex_no_writeoff_refund_transfer_beta")
    current_payload = {key: b"workbook" for key in keys}
    current_payload["beta_export_variant"] = app_workflows.BETA_EXPORT_VARIANT
    stale_payload = dict(current_payload, beta_export_variant="beta_market_ecommerce_peer_sales_points_v4")
    incomplete_payload = dict(current_payload)
    incomplete_payload.pop("ex_no_writeoff_beta")

    assert app_workflows._beta_export_cache_is_current(current_payload)
    assert not app_workflows._beta_export_cache_is_current(stale_payload)
    assert not app_workflows._beta_export_cache_is_current(incomplete_payload)


def test_beta_cache_identity_differs_from_formal_identity():
    import app_workflows

    assert app_workflows._build_export_variant_key(
        "source", "rules", "beta_ecommerce_sales_point_v1"
    ) != app_workflows._build_export_variant_key("source", "rules", "official")


def test_export_ui_exposes_beta_ecommerce_comparison_without_replacing_official_action():
    source = Path("app_pages.py").read_text(encoding="utf-8")

    assert "Beta comparison export" in source
    assert BETA_LABEL in source
    assert "ex_no_writeoff_refund_transfer" in source
